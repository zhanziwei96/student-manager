"""
学生API控制器 - FastAPI版本
"""
from typing import List, Optional
from fastapi import APIRouter, Depends, Request, HTTPException, Query, UploadFile, File
from pydantic import BaseModel, Field

from infrastructure.persistence.database import Database
from infrastructure.persistence.repositories.sqlite_student_repository import SQLiteStudentRepository
from infrastructure.persistence.repositories.sqlite_score_log_repository import SQLiteScoreLogRepository
from infrastructure.security.rate_limiter import score_limit
from infrastructure.security.session import require_login, is_admin, get_session_user_id
from infrastructure.security.xss_protection import sanitize_student_list, sanitize_student_data
from infrastructure.logging import logger
from infrastructure.cache import get_cache_client, CacheKeyBuilder, evict_student_cache
from infrastructure.config import AuthConfig, CacheConfig, HttpStatus, ScoreConfig, ImportExportConfig
from application.services.student_app_service import StudentAppService
from application.dto.student_dto import CreateStudentDTO, UpdateScoreDTO


router = APIRouter(prefix="/api", tags=["students"])


# ============ Pydantic模型 ============

class CreateStudentRequest(BaseModel):
    student_id: str = Field(..., min_length=AuthConfig.NAME_MIN_LENGTH, description="学号")
    name: str = Field(..., min_length=AuthConfig.NAME_MIN_LENGTH, description="姓名")
    class_name: Optional[str] = Field(None, description="班级")


class UpdateScoreRequest(BaseModel):
    score_change: float = Field(..., description="分数变动值")
    reason: str = Field(..., min_length=AuthConfig.NAME_MIN_LENGTH, description="变动原因")


class ApiResponse(BaseModel):
    success: bool
    message: Optional[str] = None
    data: Optional[dict] = None


# ============ 依赖注入 ============

def get_student_service():
    """获取学生应用服务"""
    db = Database()
    repo = SQLiteStudentRepository(db)
    log_repo = SQLiteScoreLogRepository(db)
    return StudentAppService(repo, log_repo)


def get_current_user_id(request: Request) -> int:
    """获取当前用户ID（依赖注入）"""
    return require_login(request)


# ============ API端点 ============

@router.get("/students", response_model=dict)
async def get_students(
    request: Request,
    class_name: Optional[str] = Query(None, description="班级名称"),
    service: StudentAppService = Depends(get_student_service),
    nocache: bool = Query(False, description="跳过缓存")
):
    """获取学生列表"""
    require_login(request)
    
    # 尝试从缓存获取
    cache = get_cache_client()
    if not nocache and cache.is_connected:
        cache_key = CacheKeyBuilder.student_list(class_name)
        cached_data = await cache.get_json(cache_key)
        if cached_data is not None:
            logger.debug(f"Student list cache hit: {cache_key}")
            return {'success': True, 'data': cached_data}
    
    if class_name:
        students = service.get_students_by_class(class_name)
    else:
        students = service.get_all_students()
    
    # XSS防护：对学生数据进行HTML转义
    student_data = [sanitize_student_data(s.__dict__) for s in students]
    
    # 写入缓存
    if not nocache and cache.is_connected:
        await cache.set_json(cache_key, student_data, ttl=CacheConfig.STUDENT_LIST_TTL_SECONDS)
    
    return {
        'success': True,
        'data': student_data
    }


@router.get("/students/{student_id}", response_model=dict)
async def get_student(
    request: Request,
    student_id: str,
    service: StudentAppService = Depends(get_student_service),
    nocache: bool = Query(False, description="跳过缓存")
):
    """获取单个学生"""
    require_login(request)
    
    # 尝试从缓存获取
    cache = get_cache_client()
    if not nocache and cache.is_connected:
        cache_key = CacheKeyBuilder.student_detail(student_id)
        cached_data = await cache.get_json(cache_key)
        if cached_data is not None:
            logger.debug(f"Student detail cache hit: {cache_key}")
            return {'success': True, 'data': cached_data}
    
    student = service.get_student_by_id(student_id)
    if student:
        # XSS防护：对学生数据进行HTML转义
        result = sanitize_student_data(student.__dict__)
        
        # 写入缓存
        if not nocache and cache.is_connected:
            cache_key = CacheKeyBuilder.student_detail(student_id)
            await cache.set_json(cache_key, result, ttl=CacheConfig.STUDENT_DETAIL_TTL_SECONDS)
        
        return {
            'success': True,
            'data': result
        }
    raise HTTPException(status_code=HttpStatus.NOT_FOUND, detail='学生不存在')


@router.post("/students", response_model=dict)
async def create_student(
    request: Request,
    data: CreateStudentRequest,
    service: StudentAppService = Depends(get_student_service)
):
    """创建学生"""
    require_login(request)
    
    try:
        dto = CreateStudentDTO(
            student_id=data.student_id.strip(),
            name=data.name.strip(),
            class_name=data.class_name.strip() if data.class_name else None
        )
        student = service.create_student(dto)
        
        # 清除相关缓存
        cache = get_cache_client()
        if cache.is_connected:
            # 清除学生列表缓存
            await cache.delete(CacheKeyBuilder.student_list())
            await cache.delete(CacheKeyBuilder.student_list(data.class_name))
            # 清除统计缓存
            await cache.delete(CacheKeyBuilder.student_stats())
            logger.debug("Student cache evicted after create")
        
        return {
            'success': True,
            'message': '学生创建成功',
            'data': student.__dict__
        }
    except ValueError as e:
        raise HTTPException(status_code=HttpStatus.BAD_REQUEST, detail=str(e))


@router.post(
    "/students/{student_id}/score", 
    response_model=dict,
    dependencies=[Depends(score_limit())]
)
async def update_score(
    request: Request,
    student_id: str,
    data: UpdateScoreRequest,
    service: StudentAppService = Depends(get_student_service)
):
    """更新学生分数"""
    require_login(request)
    
    try:
        dto = UpdateScoreDTO(
            student_id=student_id,
            delta=float(data.score_change),
            reason=data.reason.strip(),
            operator=request.session.get('username', 'unknown')
        )
        student = service.update_score(dto)
        
        # 清除相关缓存
        cache = get_cache_client()
        if cache.is_connected:
            # 清除该学生详情缓存
            await cache.delete(CacheKeyBuilder.student_detail(student_id))
            # 清除学生列表缓存
            await cache.delete(CacheKeyBuilder.student_list())
            if student and student.class_name:
                await cache.delete(CacheKeyBuilder.student_list(student.class_name))
            # 清除统计缓存
            await cache.delete(CacheKeyBuilder.student_stats())
            logger.debug(f"Student cache evicted after score update: {student_id}")
        
        return {
            'success': True,
            'message': '分数更新成功',
            'data': student.__dict__
        }
    except ValueError as e:
        raise HTTPException(status_code=HttpStatus.BAD_REQUEST, detail=str(e))


@router.delete("/students/{student_id}", response_model=dict)
async def delete_student(
    request: Request,
    student_id: str,
    service: StudentAppService = Depends(get_student_service)
):
    """删除学生"""
    require_login(request)
    
    try:
        # 先获取学生信息（用于清除缓存）
        student = service.get_student_by_id(student_id)
        class_name = student.class_name if student else None
        
        service.delete_student(student_id)
        
        # 清除相关缓存
        cache = get_cache_client()
        if cache.is_connected:
            # 清除该学生详情缓存
            await cache.delete(CacheKeyBuilder.student_detail(student_id))
            # 清除学生列表缓存
            await cache.delete(CacheKeyBuilder.student_list())
            if class_name:
                await cache.delete(CacheKeyBuilder.student_list(class_name))
            # 清除统计缓存
            await cache.delete(CacheKeyBuilder.student_stats())
            logger.debug(f"Student cache evicted after delete: {student_id}")
        
        return {'success': True, 'message': '学生删除成功'}
    except Exception as e:
        logger.error(f"删除学生失败: {e}", exc_info=True)
        raise HTTPException(status_code=HttpStatus.INTERNAL_ERROR, detail=str(e))


@router.get("/score/logs", response_model=dict)
async def get_score_logs(
    request: Request,
    student_id: Optional[str] = Query(None),
    class_name: Optional[str] = Query(None)
):
    """获取分数变更日志"""
    require_login(request)
    
    db = Database()
    log_repo = SQLiteScoreLogRepository(db)
    
    logs = log_repo.find_by_filters(
        student_id=student_id,
        class_name=class_name,
        limit=PaginationConfig.MAX_SCORE_LOGS
    )
    
    return {
        'success': True,
        'data': [log.to_dict() for log in logs]
    }


@router.get("/stats", response_model=dict)
async def get_stats(
    request: Request,
    nocache: bool = Query(False, description="跳过缓存")
):
    """获取统计数据"""
    require_login(request)
    
    # 尝试从缓存获取
    cache = get_cache_client()
    if not nocache and cache.is_connected:
        cache_key = CacheKeyBuilder.student_stats()
        cached_data = await cache.get_json(cache_key)
        if cached_data is not None:
            logger.debug(f"Stats cache hit: {cache_key}")
            return {'success': True, 'data': cached_data}
    
    db = Database()
    
    with db.connection() as conn:
        # 总学生数
        cursor = conn.execute("SELECT COUNT(*) as count FROM students")
        total_students = cursor.fetchone()['count']
        
        # 班级列表
        cursor = conn.execute(
            "SELECT class_name, COUNT(*) as count FROM students GROUP BY class_name"
        )
        class_stats = [
            {'class_name': row['class_name'], 'student_count': row['count']}
            for row in cursor.fetchall()
        ]
        
        # 今日签到人数
        from datetime import datetime
        today = datetime.now().strftime('%Y-%m-%d')
        cursor = conn.execute(
            "SELECT COUNT(DISTINCT student_id) as count FROM checkin_records WHERE checkin_date = ?",
            (today,)
        )
        today_checkins = cursor.fetchone()['count']
        
        # 平均分
        cursor = conn.execute("SELECT AVG(score) as avg FROM students")
        avg_score = cursor.fetchone()['avg'] or 0
    
    result = {
        'total_students': total_students,
        'class_count': len(class_stats),
        'class_stats': class_stats,
        'today_checkins': today_checkins,
        'average_score': round(avg_score, 2)  # 保留两位小数
    }
    
    # 写入缓存
    if not nocache and cache.is_connected:
        await cache.set_json(CacheKeyBuilder.student_stats(), result, ttl=CacheConfig.STATS_TTL_SECONDS)
    
    return {
        'success': True,
        'data': result
    }


@router.post("/admin/reset-scores", response_model=dict)
async def reset_all_scores(
    request: Request,
    service: StudentAppService = Depends(get_student_service)
):
    """重置所有学生分数（管理员）"""
    require_admin(request)
    
    from domain.value_objects.score import Score
    db = Database()
    repo = SQLiteStudentRepository(db)
    
    students = repo.find_all()
    for student in students:
        student.score = Score(ScoreConfig.DEFAULT_SCORE)
        repo.save(student)
    
    # 清除所有学生相关缓存
    cache = get_cache_client()
    if cache.is_connected:
        count = await cache.delete_pattern(CacheKeyBuilder.pattern_student_all())
        logger.info(f"Cache cleared after reset-scores: {count} entries")
    
    return {
        'success': True,
        'message': f'已重置{len(students)}名学生的分数'
    }


@router.post("/students/import", response_model=dict)
async def import_students(
    request: Request,
    file: UploadFile = File(...),
    service: StudentAppService = Depends(get_student_service)
):
    """从Excel导入学生"""
    require_login(request)
    
    try:
        import io
        from openpyxl import load_workbook
        
        if not file.filename.endswith(('.xlsx', '.xls')):
            raise HTTPException(status_code=HttpStatus.BAD_REQUEST, detail='请上传Excel文件')
        
        # 读取Excel
        contents = await file.read()
        stream = io.BytesIO(contents)
        wb = load_workbook(stream)
        ws = wb.active
        
        success_count = 0
        error_count = 0
        errors = []
        
        # 从第二行开始读取（第一行为表头）
        for row_num, row in enumerate(ws.iter_rows(min_row=ImportExportConfig.EXCEL_START_ROW, values_only=True), start=ImportExportConfig.EXCEL_START_ROW):
            try:
                if not row or len(row) < 2:
                    continue
                
                student_id = str(row[0]).strip() if row[0] else None
                name = str(row[1]).strip() if row[1] else None
                class_name = str(row[2]).strip() if len(row) > 2 and row[2] else None
                
                if not student_id or not name:
                    continue
                
                dto = CreateStudentDTO(
                    student_id=student_id,
                    name=name,
                    class_name=class_name
                )
                service.create_student(dto)
                success_count += 1
                
            except ValueError as e:
                error_count += 1
                errors.append(f"第{row_num}行: {str(e)}")
        
        # 清除学生相关缓存
        cache = get_cache_client()
        if cache.is_connected and success_count > 0:
            count = await cache.delete_pattern(CacheKeyBuilder.pattern_student_all())
            logger.info(f"Cache cleared after import: {count} entries")
        
        return {
            'success': True,
            'message': f'导入完成: 成功{success_count}条, 失败{error_count}条',
            'data': {
                'success_count': success_count,
                'error_count': error_count,
                'errors': errors[:ImportExportConfig.MAX_IMPORT_ERRORS_DISPLAY]
            }
        }
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"导入学生失败: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f'导入失败: {str(e)}')
